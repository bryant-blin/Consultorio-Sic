from flask import Flask, render_template, request, redirect, url_for, send_file, flash, jsonify
import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from flask_login import LoginManager, login_required, current_user
from functools import wraps
from modelo.usuarios import db, Login, Usuario, Rol, Historial_Medico, Cita, Examen, Categoria, Factura, PagoDetalle, Configuracion
from controladores.autenticacion import auth_login, auth_logout
from datetime import datetime, date, time
import os
import sys
import urllib.parse
try:
    import webview
except ImportError:
    webview = None
import uuid
import requests
from bs4 import BeautifulSoup
from threading import Thread
from openpyxl.drawing.image import Image
from sqlalchemy import func
from openpyxl.drawing.image import Image
from sqlalchemy import func

# Importar el servicio del Bot de Telegram
from servicios.bot_telegram import start_bot_thread, notificar_staff, notificar_paciente

def roles_required(*roles):
    """Decorador para proteger rutas según el rol del usuario actual."""
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if not current_user.is_authenticated:
                return redirect(url_for('login'))
            if current_user.rol_perfil.nombre not in roles:
                flash("No tienes permisos para acceder a esta función.", "error")
                return redirect(url_for('dashboard'))
            return f(*args, **kwargs)
        return decorated_function
    return decorator
def resource_path(relative_path):
    """ Obtiene la ruta absoluta para recursos, funciona para dev y PyInstaller """
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)


app = Flask(__name__,
            template_folder=resource_path('vista'),
            static_folder=resource_path('static'))

# ═══════════════════════════════════════════════════════════
# CONFIGURACIÓN DE BASE DE DATOS Y ENTORNO
# ═══════════════════════════════════════════════════════════

# Prioridad 1: Variable de entorno (Nube) | Prioridad 2: Configuración Local
database_url = os.environ.get('DATABASE_URL')

if database_url:
    # Ajuste para Render/Heroku que usan postgres:// en lugar de postgresql://
    if database_url.startswith("postgres://"):
        database_url = database_url.replace("postgres://", "postgresql://", 1)
    app.config['SQLALCHEMY_DATABASE_URI'] = database_url
else:
    # Configuración local manual
    user = 'postgres'
    password = urllib.parse.quote_plus('0111')
    host = 'localhost'
    database = 'data_sic'
    app.config['SQLALCHEMY_DATABASE_URI'] = f'postgresql://{user}:{password}@{host}:5432/{database}?client_encoding=utf8'

app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'secreto_clinica_sic')

db.init_app(app)
login_manager = LoginManager(app)
login_manager.login_view = 'auth_login'
login_manager.login_message = "Por favor, inicie sesión para acceder a esta página."


@login_manager.user_loader
def load_user(id_login):
    return Login.query.get(int(id_login))


def seed_roles():
    """Crea los roles básicos si no existen"""
    roles_basicos = ['Administrador', 'Secretaria', 'Doctor']
    for nombre_rol in roles_basicos:
        rol_existente = Rol.query.filter_by(nombre=nombre_rol).first()
        if not rol_existente:
            nuevo_rol = Rol(nombre=nombre_rol)
            db.session.add(nuevo_rol)
    db.session.commit()


# ═══════════════════════════════════════════════════════════
# INICIALIZACIÓN DE LA APLICACIÓN
# ═══════════════════════════════════════════════════════════

def inicializar_todo():
    with app.app_context():
        try:
            print("[SISTEMA]: Verificando conexión a Base de Datos...")
            db.create_all()
            print("[SISTEMA]: Base de Datos lista.")
            
            print("[SISTEMA]: Inicializando roles...")
            seed_roles()
            
            print("[SISTEMA]: Iniciando hilo del Bot de Telegram...")
            start_bot_thread(app, db, Cita, Historial_Medico)
            print("[SISTEMA]: Bot de Telegram en ejecución.")
            
        except Exception as e:
            print(f"[ERROR]: Fallo en la inicialización: {e}")

# Llamar a la inicialización antes de que el servidor atienda peticiones
inicializar_todo()


# ═══════════════════════════════════════════════════════════
# LÓGICA DE TASA BCV (SCRAPING)
# ═══════════════════════════════════════════════════════════

def obtener_tasa_bcv():
    """Obtiene la tasa oficial del Dólar desde la web del BCV."""
    try:
        url = "https://www.bcv.org.ve/"
        headers = {'User-Agent': 'Mozilla/5.0'}
        response = requests.get(url, headers=headers, timeout=10, verify=False)
        soup = BeautifulSoup(response.text, 'html.parser')
        
        # El BCV tiene el dolar en un div con id 'dolar'
        tasa_str = soup.find('div', id='dolar').find('strong').get_text().strip()
        tasa_val = float(tasa_str.replace(',', '.'))
        
        # Guardar en base de datos
        with app.app_context():
            config = Configuracion.query.first()
            if not config:
                config = Configuracion(tasa_bcv=tasa_val)
                db.session.add(config)
            else:
                config.tasa_bcv = tasa_val
                config.ultima_actualizacion = datetime.now()
            db.session.commit()
        return tasa_val
    except Exception as e:
        print(f"Error al obtener tasa BCV: {e}")
        # Intentar retornar la ultima guardada
        with app.app_context():
            config = Configuracion.query.first()
            return float(config.tasa_bcv) if config else 0.0

@app.route('/tasa/actualizar')
@login_required
def actualizar_tasa_api():
    """Endpoint para forzar el refresco de la tasa desde el frontend."""
    nueva_tasa = obtener_tasa_bcv()
    return jsonify({"tasa": float(nueva_tasa)})

@app.context_processor
def inject_tasa():
    """Inyectar la tasa BCV en todos los templates Jinja2."""
    config = Configuracion.query.first()
    tasa = float(config.tasa_bcv) if config else 1.0
    return dict(tasa_bcv=tasa)

# ═══════════════════════════════════════════════════════════
# RUTAS DE AUTENTICACIÓN
# ═══════════════════════════════════════════════════════════
app.add_url_rule('/login', view_func=auth_login, methods=['GET', 'POST'])
app.add_url_rule('/logout', view_func=auth_logout)


@app.route('/registro', methods=['GET', 'POST'])
def registro():
    from controladores.autenticacion import auth_registro
    return auth_registro()


@app.route('/recuperar', methods=['GET', 'POST'])
def recuperar():
    from controladores.autenticacion import auth_recuperar
    return auth_recuperar()


@app.route('/')
def inicio():
    return redirect(url_for('auth_login'))


# ═══════════════════════════════════════════════════════════
# DASHBOARD PRINCIPAL
# ═══════════════════════════════════════════════════════════
@app.route('/dashboard')
@login_required
def dashboard():
    # Contadores para las tarjetas del dashboard
    total_citas = Cita.query.count()
    total_examenes = Examen.query.count()
    total_facturas = Factura.query.filter_by(estado_cierre=False).count()
    total_pacientes = Historial_Medico.query.count()
    return render_template('dashboard.html',
                           total_citas=total_citas,
                           total_examenes=total_examenes,
                           total_facturas=total_facturas,
                           total_pacientes=total_pacientes)


# ═══════════════════════════════════════════════════════════
# PUNTO DE VENTA (POS)
# ═══════════════════════════════════════════════════════════
@app.route('/pos')
@login_required
@roles_required('Administrador', 'Secretaria')
def punto_de_venta():
    # Capturamos el nombre si viene de la lista de pacientes
    paciente_seleccionado = request.args.get('paciente', '') 
    examenes = Examen.query.order_by(Examen.nombre_examen.asc()).all()
    categorias = Categoria.query.order_by(Categoria.nombre.asc()).all()
    return render_template('panel_principal.html', 
                           examenes=examenes, 
                           categorias=categorias, 
                           paciente_previo=paciente_seleccionado) # Pasamos el nombre



@app.route('/pos/validar', methods=['POST'])
@login_required
@roles_required('Administrador', 'Secretaria')
def pos_validar():
    data = request.json
    try:
        # Generar un ID unico para esta transaccion/venta
        id_trans = str(uuid.uuid4())[:8].upper()
        
        # Guardar los items (Examenes)
        for item in data['items']:
            # Repetir la insercion segun la cantidad vendida
            qty = int(item.get('cantidad', 1))
            for _ in range(qty):
                nueva = Factura(
                    nombre_paciente=data['paciente'] or "Consumidor Final",
                    id_examenes=item['id'],
                    precio_final=item['precio'],
                    id_transaccion=id_trans
                )
                db.session.add(nueva)
        
        # Guardar los pagos realizados (Pagos Mixtos)
        if 'pagos' in data:
            for pago in data['pagos']:
                nuevo_pago = PagoDetalle(
                    id_transaccion=id_trans,
                    metodo=pago['metodo'],
                    monto=pago['monto']
                )
                db.session.add(nuevo_pago)
        
        db.session.commit()
        return {"status": "success", "message": "Venta registrada con éxito", "id_venta": id_trans}, 200
    except Exception as e:
        db.session.rollback()
        print(f"Error en validación POS: {e}")
        return {"status": "error", "message": str(e)}, 500


# ═══════════════════════════════════════════════════════════
# CRUD DE PACIENTES
# ═══════════════════════════════════════════════════════════
@app.route('/pacientes')
@login_required
def lista_pacientes():
    query = request.args.get('query', '')
    modo_seleccion = request.args.get('select', 'false')
    
    if query:
        # Busca por cédula o nombre
        pacientes = Historial_Medico.query.filter(
            (Historial_Medico.cedula.like(f'%{query}%')) |
            (Historial_Medico.nombre_paciente.like(f'%{query}%'))
        ).all()
    else:
        pacientes = Historial_Medico.query.order_by(Historial_Medico.nombre_paciente.asc()).all()
        
    return render_template('pacientes.html', pacientes=pacientes, query=query, select=modo_seleccion)


@app.route('/pacientes/nuevo', methods=['POST'])
@login_required
@roles_required('Administrador', 'Secretaria')
def nuevo_paciente_crud():
    try:
        nuevo = Historial_Medico(
            cedula=request.form.get('cedula'),
            nombre_paciente=f"{request.form.get('nombre')} {request.form.get('apellido')}",
            edad=request.form.get('edad'),
            telefono=request.form.get('telefono')
        )
        db.session.add(nuevo)
        db.session.commit()
        flash("Paciente registrado con éxito.", "success")
    except Exception as e:
        db.session.rollback()
        flash(f"Error: {str(e)}", "error")
    return redirect(url_for('lista_pacientes'))


@app.route('/pacientes/editar/<int:id>', methods=['POST'])
@login_required
def editar_paciente_crud(id):
    p = Historial_Medico.query.get_or_404(id)
    p.cedula = request.form.get('cedula')
    p.nombre_paciente = request.form.get('nombre_completo')
    p.telefono = request.form.get('telefono')
    p.edad = request.form.get('edad')
    db.session.commit()
    flash("Datos actualizados.", "success")
    return redirect(url_for('lista_pacientes'))


@app.route('/pacientes/eliminar/<int:id>')
@login_required
def eliminar_paciente_crud(id):
    p = Historial_Medico.query.get_or_404(id)
    db.session.delete(p)
    db.session.commit()
    flash("Paciente eliminado.", "success")
    return redirect(url_for('lista_pacientes'))


# ═══════════════════════════════════════════════════════════
# CONSULTAS / EXÁMENES
# ═══════════════════════════════════════════════════════════
@app.route('/consultas')
@login_required
@roles_required('Administrador', 'Secretaria')
def consultas():
    examenes = Examen.query.order_by(Examen.nombre_examen.asc()).all()
    categorias = Categoria.query.order_by(Categoria.nombre.asc()).all()
    return render_template('consultas.html', examenes=examenes, categorias=categorias)


@app.route('/consultas/nueva', methods=['POST'])
@login_required
@roles_required('Administrador', 'Secretaria')
def nueva_consulta():
    config = Configuracion.query.first()
    tasa = float(config.tasa_bcv) if config else 1.0
    precio_usd = float(request.form.get('precio_usd', 0))
    
    nuevo = Examen(
        nombre_examen=request.form.get('nombre'),
        id_categoria=request.form.get('id_categoria'),
        precio=precio_usd,  # Guardar directamente en USD
        descripcion=request.form.get('descripcion')
    )
    db.session.add(nuevo)
    db.session.commit()
    return redirect(url_for('consultas'))


@app.route('/consultas/eliminar/<int:id>')
@login_required
@roles_required('Administrador', 'Secretaria')
def eliminar_consulta(id):
    examen = Examen.query.get_or_404(id)
    db.session.delete(examen)
    db.session.commit()
    return redirect(url_for('consultas'))


@app.route('/consultas/editar/<int:id>', methods=['POST'])
@login_required
@roles_required('Administrador', 'Secretaria')
def editar_consulta(id):
    config = Configuracion.query.first()
    tasa = float(config.tasa_bcv) if config else 1.0
    precio_usd = float(request.form.get('precio_usd', 0))
    
    examen = Examen.query.get_or_404(id)
    examen.nombre_examen = request.form.get('nombre')
    examen.id_categoria = request.form.get('id_categoria')
    examen.precio = precio_usd  # Guardar directamente en USD
    examen.descripcion = request.form.get('descripcion')
    db.session.commit()
    return redirect(url_for('consultas'))


# ═══════════════════════════════════════════════════════════
# REPORTES / FACTURACIÓN
# ═══════════════════════════════════════════════════════════
@app.route('/reportes')
@login_required
@roles_required('Administrador', 'Secretaria')
def reportes():
    # Solo mostrar facturas que NO han sido cerradas (estado_cierre=False)
    facturas_all = Factura.query.filter_by(estado_cierre=False).order_by(Factura.fecha_facturacion.desc()).all()
    
    # Agrupar en Python para mayor control
    transacciones = {}
    for f in facturas_all:
        tid = f.id_transaccion if f.id_transaccion else str(f.id_factura)
        
        if tid not in transacciones:
            pagos = []
            if f.id_transaccion:
                pagos = PagoDetalle.query.filter_by(id_transaccion=f.id_transaccion).all()
            
            transacciones[tid] = {
                'id_factura_referencia': f.id_factura,
                'id_transaccion': f.id_transaccion,
                'fecha': f.fecha_facturacion,
                'paciente': f.nombre_paciente,
                'detalles_conteo': {}, # Diccionario temporal para agrupar nombres
                'total': 0.0,
                'pagos': pagos
            }
        
        nombre_ex = f.examen.nombre_examen if f.examen else "Desconocido"
        # Contar cantidades del mismo examen
        transacciones[tid]['detalles_conteo'][nombre_ex] = transacciones[tid]['detalles_conteo'].get(nombre_ex, 0) + 1
        transacciones[tid]['total'] += float(f.precio_final)

    # Formatear los nombres de los exámenes con sus cantidades
    for tid in transacciones:
        conteo = transacciones[tid]['detalles_conteo']
        resumen = []
        for nombre, qty in conteo.items():
            if qty > 1:
                resumen.append(f"{nombre} (x{qty})")
            else:
                resumen.append(nombre)
        transacciones[tid]['examenes'] = resumen

    # Convertir a lista y ordenar
    facturas_agrupadas = sorted(transacciones.values(), key=lambda x: x['fecha'], reverse=True)
    
    examenes = Examen.query.order_by(Examen.nombre_examen.asc()).all()
    return render_template('reportes.html', facturas=facturas_agrupadas, examenes=examenes)


@app.route('/reportes/eliminar/<int:id>')
@login_required
@roles_required('Administrador', 'Secretaria')
def eliminar_factura(id):
    fac = Factura.query.get_or_404(id)
    try:
        if fac.id_transaccion:
            # ANULAR VENTA COMPLETA: Borrar todos los items y pagos de esa transaccion
            Factura.query.filter_by(id_transaccion=fac.id_transaccion).delete()
            PagoDetalle.query.filter_by(id_transaccion=fac.id_transaccion).delete()
        else:
            # Compatibilidad: Borrar solo el registro individual antiguo
            db.session.delete(fac)
        
        db.session.commit()
        flash("Transacción anulada con éxito.", "success")
    except Exception as e:
        db.session.rollback()
        flash(f"Error al anular: {str(e)}", "error")
    return redirect(url_for('reportes'))


@app.route('/reportes/editar/<int:id>', methods=['POST'])
@login_required
@roles_required('Administrador', 'Secretaria')
def editar_factura(id):
    fac = Factura.query.get_or_404(id)
    fac.nombre_paciente = request.form.get('nombre_paciente')
    fac.id_examenes = request.form.get('id_examen')
    fac.precio_final = request.form.get('precio_final')
    db.session.commit()
    return redirect(url_for('reportes'))


@app.route('/reportes/exportar')
@login_required
@roles_required('Administrador', 'Secretaria')
def exportar_reportes():
    # Exportar solo las facturas activas del cierre actual
    facturas_all = Factura.query.filter_by(estado_cierre=False).order_by(Factura.fecha_facturacion.desc()).all()
    
    transacciones = {}
    for f in facturas_all:
        tid = f.id_transaccion if f.id_transaccion else str(f.id_factura)
        if tid not in transacciones:
            metodos_texto = "Registro Antiguo"
            if f.id_transaccion:
                pagos = PagoDetalle.query.filter_by(id_transaccion=f.id_transaccion).all()
                metodos_texto = ", ".join([f"{p.metodo} (Bs. {p.monto})" for p in pagos])
            
            transacciones[tid] = {
                'fecha': f.fecha_facturacion,
                'paciente': f.nombre_paciente,
                'detalles_conteo': {},
                'total': 0.0,
                'metodos': metodos_texto
            }
        
        nombre_ex = f.examen.nombre_examen if f.examen else "Desconocido"
        transacciones[tid]['detalles_conteo'][nombre_ex] = transacciones[tid]['detalles_conteo'].get(nombre_ex, 0) + 1
        transacciones[tid]['total'] += float(f.precio_final)

    for tid in transacciones:
        conteo = transacciones[tid]['detalles_conteo']
        resumen = [f"{n} (x{q})" if q > 1 else n for n, q in conteo.items()]
        transacciones[tid]['examenes_final'] = ", ".join(resumen)

    facturas_agrupadas = sorted(transacciones.values(), key=lambda x: x['fecha'], reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Cierre de Ventas"
    
    # Estilos Premium (Estilo Verde similar a imagen de referencia)
    verde_oscuro = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid") # Verde esmeralda para el título principal
    verde_claro = PatternFill(start_color="A7F3D0", end_color="A7F3D0", fill_type="solid") # Verde más suave para cabeceras de columnas
    verde_fila_par = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid") # Verde super claro para filas
    blanco = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    fuente_titulo = Font(color="FFFFFF", bold=True, size=16)
    fuente_cabecera = Font(color="064E3B", bold=True, size=12) # Texto verde muy oscuro/casí negro
    fuente_normal = Font(color="1E293B", size=11)
    
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    borde_fino = Border(left=Side(style='thin', color='B8D8C9'), 
                        right=Side(style='thin', color='B8D8C9'), 
                        top=Side(style='thin', color='B8D8C9'), 
                        bottom=Side(style='thin', color='B8D8C9'))

    # Fila de Título Principal (Cierre de Ventas)
    ws.append(["Resumen de Ventas", "", "", "", ""])
    ws.merge_cells("A1:E1")
    for cell in ws[1]:
        cell.fill = verde_oscuro
        cell.font = fuente_titulo
        cell.alignment = centro
        cell.border = borde_fino
    ws.row_dimensions[1].height = 30

    # Cabeceras de Columna
    headers = ["Fecha", "Paciente", "Detalle de Venta", "Métodos de Pago", "Total (Bs.)"]
    ws.append(headers)
    for col_idx, cell in enumerate(ws[2], 1):
        cell.fill = verde_claro
        cell.font = fuente_cabecera
        cell.alignment = centro
        cell.border = borde_fino

    # Insertar Datos
    start_row = 3
    for idx, t in enumerate(facturas_agrupadas):
        ws.append([
            t['fecha'].strftime('%d/%m/%Y %H:%M'), 
            t['paciente'], 
            t['examenes_final'], 
            t['metodos'], 
            float(t['total'])
        ])
        
        # Aplicar estilo a las filas
        fill_color = verde_fila_par if idx % 2 == 1 else blanco
        for cell in ws[start_row + idx]:
            cell.fill = fill_color
            cell.font = fuente_normal
            cell.alignment = centro
            cell.border = borde_fino
            
            # Formato de moneda para última columna
            if cell.column == 5:
                cell.number_format = 'Bs. #,##0.00'

    # Ajustar Ancho de Columnas
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 35
    ws.column_dimensions['D'].width = 35
    ws.column_dimensions['E'].width = 18
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"Cierre_Ventas_{datetime.now().strftime('%d_%m_%Y')}.xlsx"
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=filename)


@app.route('/reportes/finalizar-dia', methods=['POST'])
@login_required
@roles_required('Administrador', 'Secretaria')
def finalizar_dia_api():
    """Marca todas las ventas y pacientes nuevos actuales como finalizados (Cierre de Día)."""
    try:
        # Marcar todas las facturas no cerradas como cerradas
        Factura.query.filter_by(estado_cierre=False).update({Factura.estado_cierre: True})
        
        # Marcar todos los pacientes registrados hoy como "ya procesados" en este cierre
        Historial_Medico.query.filter_by(estado_cierre=False).update({Historial_Medico.estado_cierre: True})
        
        db.session.commit()
        return jsonify({"status": "success", "message": "Día finalizado con éxito"})
    except Exception as e:
        db.session.rollback()
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route('/api/stats')
@login_required
def api_stats():
    """Analítica BI avanzada: Financiera, Operativa y Fidelización (Hoy, Mes, Año)."""
    import calendar
    periodo = request.args.get('periodo', 'hoy')
    hoy = date.today()
    
    # 0. Datos de Configuración (Tasa BCV)
    config = Configuracion.query.first()
    tasa = float(config.tasa_bcv) if config else 1.0

    filtros = []
    filtros_pacientes = []
    
    if periodo == 'hoy':
        inicio_dia = datetime.combine(hoy, time.min)
        fin_dia = datetime.combine(hoy, time.max)
        grafica_agrupacion = func.extract('hour', Factura.fecha_facturacion)
        # Hoy: Solo ventas NO cerradas
        filtros.append(Factura.estado_cierre == False)
        filtros_pacientes.append(Historial_Medico.estado_cierre == False)
    elif periodo == 'mes':
        inicio_dia = datetime(hoy.year, hoy.month, 1, 0, 0, 0)
        ultimo_dia = calendar.monthrange(hoy.year, hoy.month)[1]
        fin_dia = datetime(hoy.year, hoy.month, ultimo_dia, 23, 59, 59)
        grafica_agrupacion = func.extract('day', Factura.fecha_facturacion) # agrupar por dia
    elif periodo == 'anio':
        inicio_dia = datetime(hoy.year, 1, 1, 0, 0, 0)
        fin_dia = datetime(hoy.year, 12, 31, 23, 59, 59)
        grafica_agrupacion = func.extract('month', Factura.fecha_facturacion) # agrupar por mes
        
    filtros.append(Factura.fecha_facturacion.between(inicio_dia, fin_dia))
    filtros_pacientes.append(Historial_Medico.fecha_registro.between(inicio_dia, fin_dia))

    # 1. KPIs Financieros
    query_base = Factura.query.filter(*filtros)
    total_bs = db.session.query(func.sum(Factura.precio_final)).filter(*filtros).scalar() or 0
    cantidad_ventas = query_base.count()
    ticket_promedio = float(total_bs) / cantidad_ventas if cantidad_ventas > 0 else 0
    
    # 2. Fidelización: PACIENTES NUEVOS REGISTRADOS
    pacientes_nuevos_count = Historial_Medico.query.filter(*filtros_pacientes).count()
    
    # Recurrentes (Ventas del periodo de pacientes que ya existían antes - o que compran varias veces)
    subquery_previos = db.session.query(Factura.nombre_paciente).filter(Factura.fecha_facturacion < inicio_dia).subquery()
    ventas_nuevos_historico = db.session.query(func.count(func.distinct(Factura.nombre_paciente)))\
        .filter(*filtros)\
        .filter(~Factura.nombre_paciente.in_(subquery_previos)).scalar() or 0
    
    pacientes_totales_periodo = db.session.query(func.count(func.distinct(Factura.nombre_paciente)))\
        .filter(*filtros).scalar() or 0
    pacientes_recurrentes_count = pacientes_totales_periodo - ventas_nuevos_historico

    # 3. Métricas de Citas
    total_citas = Cita.query.filter(Cita.fecha_cita.between(inicio_dia, fin_dia)).count()
    citas_confirmadas = Cita.query.filter(Cita.fecha_cita.between(inicio_dia, fin_dia), Cita.estado == 'Confirmada').count()
    eficiencia_citas = (citas_confirmadas / total_citas * 100) if total_citas > 0 else 0

    # 4. Gráficas dinámicas
    # Pacientes agrupados por el periodo seleccionado (Hora, Día, o Mes)
    pacientes_cursor = db.session.query(
        grafica_agrupacion.label('agrupacion'),
        func.count(func.distinct(Factura.nombre_paciente)).label('conteo')
    ).filter(*filtros).group_by('agrupacion').all()
    v_pacientes = {int(p[0]) if p[0] else 0: p[1] for p in pacientes_cursor}

    exam_cursor = db.session.query(
        Examen.nombre_examen,
        func.count(Factura.id_factura).label('total')
    ).join(Factura, Factura.id_examenes == Examen.id_examenes)\
     .filter(*filtros)\
     .group_by(Examen.nombre_examen)\
     .order_by(func.count(Factura.id_factura).desc()).all()
    v_examenes = {n: t for n, t in exam_cursor}

    pagos_cursor = db.session.query(
        PagoDetalle.metodo,
        func.count(PagoDetalle.id_pago).label('total')
    ).join(Factura, Factura.id_transaccion == PagoDetalle.id_transaccion)\
     .filter(*filtros)\
     .group_by(PagoDetalle.metodo)\
     .order_by(func.count(PagoDetalle.id_pago).desc()).all()
    v_pagos = {m: t for m, t in pagos_cursor}

    citas_cursor = db.session.query(
        Cita.estado,
        func.count(Cita.id_cita).label('total')
    ).filter(Cita.fecha_cita.between(inicio_dia, fin_dia))\
     .group_by(Cita.estado).all()
    v_citas = {e: t for e, t in citas_cursor}

    return jsonify({
        "kpis": {
            "total_bs": float(total_bs),
            "total_usd": float(total_bs) / tasa,
            "ticket_promedio": ticket_promedio,
            "pacientes_nuevos": pacientes_nuevos_count,
            "pacientes_recurrentes": pacientes_recurrentes_count,
            "citas_confirmadas": citas_confirmadas,
            "total_citas": total_citas,
            "eficiencia_citas": eficiencia_citas,
            "tasa_actual": tasa
        },
        "pacientes": v_pacientes,
        "examenes": v_examenes,
        "pagos": v_pagos,
        "citas": v_citas
    })


# ═══════════════════════════════════════════════════════════
# HISTORIAL CLÍNICO
# ═══════════════════════════════════════════════════════════
@app.route('/historial/buscar', methods=['GET', 'POST'])
@login_required
@roles_required('Administrador', 'Doctor')
def buscar_historial():
    busqueda = request.args.get('query', '')
    resultados = []
    if busqueda:
        resultados = Historial_Medico.query.filter(
            (Historial_Medico.cedula.like(f'%{busqueda}%')) |
            (Historial_Medico.nombre_paciente.like(f'%{busqueda}%'))
        ).all()
    return render_template('historial_resultados.html', resultados=resultados, query=busqueda)


@app.route('/api/paciente/<cedula>')
@login_required
def api_buscar_paciente(cedula):
    paciente = Historial_Medico.query.filter_by(cedula=cedula).first()
    if paciente:
        return jsonify({
            'encontrado': True,
            'nombre': paciente.nombre_paciente,
            'edad': paciente.edad,
            'fecha_nac': str(paciente.fecha_nac) if paciente.fecha_nac else '',
            'telefono': paciente.telefono,
            'sexo': paciente.sexo or '',
            'direccion': paciente.direccion or '',
            'altura': paciente.altura or '',
            'peso': paciente.peso or '',
            'tension': paciente.tension or '',
            'antecedentes': paciente.antecedentes or '',
            'enfermedad_actual': paciente.enfermedad_actual or '',
            'indicaciones': paciente.indicaciones or '',
            'observaciones': paciente.observaciones or ''
        })
    return jsonify({'encontrado': False})


@app.route('/historial/ficha/<int:id>')
@login_required
@roles_required('Administrador', 'Doctor')
def ver_ficha(id):
    ficha = Historial_Medico.query.get_or_404(id)
    return render_template('historial_ficha.html', ficha=ficha, now=datetime.now())


@app.route('/historial/editar/<int:id>', methods=['GET', 'POST'])
@login_required
@roles_required('Administrador', 'Doctor')
def editar_historial(id):
    paciente = Historial_Medico.query.get_or_404(id)
    if request.method == 'POST':
        paciente.cedula = request.form.get('cedula')
        paciente.nombre_paciente = request.form.get('nombre')
        paciente.edad = request.form.get('edad')
        paciente.fecha_nac = request.form.get('fecha_nac') or None
        paciente.telefono = request.form.get('telefono')
        paciente.sexo = request.form.get('sexo')
        paciente.direccion = request.form.get('direccion')
        paciente.altura = request.form.get('altura')
        paciente.peso = request.form.get('peso')
        paciente.tension = request.form.get('tension')
        paciente.antecedentes = request.form.get('antecedentes')
        paciente.enfermedad_actual = request.form.get('enfermedad_actual')
        paciente.indicaciones = request.form.get('indicaciones')
        paciente.observaciones = request.form.get('observaciones')
        db.session.commit()
        flash("Historial actualizado con éxito.", "success")
        return redirect(url_for('ver_ficha', id=paciente.id_historial))
    return render_template('historial_nuevo.html', paciente=paciente)

# Insertar en principal.py después de ver_ficha (Línea 522 aprox.)
@app.route('/historial/visor-3d/<int:id>')
@login_required
@roles_required('Administrador', 'Doctor')
def ver_visor_3d(id):
    paciente = Historial_Medico.query.get_or_404(id)
    return render_template('visor_3d.html', p=paciente)


@app.route('/historial/nuevo', methods=['GET', 'POST'])
@login_required
@roles_required('Administrador', 'Doctor')
def nuevo_historial():
    if request.method == 'POST':
        cedula = request.form.get('cedula')
        # Buscar si ya existe para evitar UniqueViolation
        paciente = Historial_Medico.query.filter_by(cedula=cedula).first()
        
        if not paciente:
            paciente = Historial_Medico(cedula=cedula)
            db.session.add(paciente)

        paciente.nombre_paciente = request.form.get('nombre')
        paciente.edad = request.form.get('edad')
        paciente.fecha_nac = request.form.get('fecha_nac') or None
        paciente.telefono = request.form.get('telefono')
        paciente.sexo = request.form.get('sexo')
        paciente.direccion = request.form.get('direccion')
        paciente.altura = request.form.get('altura')
        paciente.peso = request.form.get('peso')
        paciente.tension = request.form.get('tension')
        paciente.antecedentes = request.form.get('antecedentes')
        paciente.enfermedad_actual = request.form.get('enfermedad_actual')
        paciente.indicaciones = request.form.get('indicaciones')
        paciente.observaciones = request.form.get('observaciones')

        db.session.commit()
        flash("Historial guardado con éxito.", "success")
        return redirect(url_for('ver_ficha', id=paciente.id_historial))
    return render_template('historial_nuevo.html')


# ═══════════════════════════════════════════════════════════
# CITAS MÉDICAS
# ═══════════════════════════════════════════════════════════
@app.route('/citas')
@login_required
def citas():
    citas = Cita.query.order_by(Cita.fecha_cita.asc()).all()
    return render_template('citas.html', citas=citas)


@app.route('/citas/nueva', methods=['POST'])
@login_required
@roles_required('Administrador', 'Secretaria')
def nueva_cita():
    try:
        fecha_str = request.form.get('fecha')
        hora_str = request.form.get('hora')
        fecha_completa = f"{fecha_str} {hora_str}"
        objeto_fecha = datetime.strptime(fecha_completa, '%Y-%m-%d %H:%M')

        # ⚠️ VALIDACIÓN DE CONFLICTO DE HORARIO
        conflicto = Cita.query.filter_by(fecha_cita=objeto_fecha).first()
        if conflicto:
            flash(f"❌ Ese horario ({hora_str}) ya está ocupado el {fecha_str}. Por favor elige otra hora.", "error")
            return redirect(url_for('citas'))

        motivo_paciente = request.form.get('motivo', '').strip()
        if not motivo_paciente:
            motivo_paciente = "Sin motivo especificado"

        nueva = Cita(
            nombre_paciente=request.form.get('nombre'),
            apellido_paciente=request.form.get('apellido'),
            cedula=request.form.get('cedula'),
            fecha_cita=objeto_fecha,
            motivo=motivo_paciente,
            estado='Pendiente'
        )
        db.session.add(nueva)
        db.session.commit()
        notificar_staff(f"📅 **Nueva Cita (vía Web)**:\n👤 {nueva.nombre_paciente} {nueva.apellido_paciente}\n🗓️ {fecha_str} a las {hora_str}\n💡 Motivo: {motivo_paciente}")
        flash("Cita agendada correctamente.", "success")
    except Exception as e:
        db.session.rollback()
        flash(f"Error al crear cita: {str(e)}", "error")
    return redirect(url_for('citas'))


@app.route('/citas/eliminar/<int:id>')
@login_required
@roles_required('Administrador', 'Secretaria')
def eliminar_cita(id):
    cita = Cita.query.get_or_404(id)
    db.session.delete(cita)
    db.session.commit()
    return redirect(url_for('citas'))


@app.route('/citas/editar/<int:id>', methods=['POST'])
@login_required
@roles_required('Administrador', 'Secretaria')
def editar_cita(id):
    try:
        cita = Cita.query.get_or_404(id)
        estado_anterior = cita.estado
        cita.nombre_paciente = request.form.get('nombre')
        cita.apellido_paciente = request.form.get('apellido')
        cita.cedula = request.form.get('cedula')
        fecha_str = request.form.get('fecha')
        hora_str = request.form.get('hora')
        fecha_completa = f"{fecha_str} {hora_str}"
        cita.fecha_cita = datetime.strptime(fecha_completa, '%Y-%m-%d %H:%M')
        cita.motivo = request.form.get('motivo')
        nuevo_estado = request.form.get('estado')
        cita.estado = nuevo_estado
        db.session.commit()

        # Notificar al paciente si el estado cambió
        if nuevo_estado != estado_anterior and cita.id_bot_sic:
            if nuevo_estado == 'Confirmada':
                notificar_paciente(cita.id_bot_sic, f"✅ *¡Tu cita ha sido enviadaaaa!*\n📅 Fecha: {fecha_str}\n⏰ Hora: {hora_str}\nrecibiras un mensaje de confirmacion pronto.")
            elif nuevo_estado == 'Cancelada':
                notificar_paciente(cita.id_bot_sic, f"❌ *Tu cita del {fecha_str} ha sido cancelada.*\nLamentamos los inconvenientes. Contáctanos para reprogramar.")

        flash("Cita actualizada con exito.", "success")
    except Exception as e:
        db.session.rollback()
        flash(f"Error al editar cita: {str(e)}", "error")
    return redirect(url_for('citas'))


@app.route('/citas/confirmar/<int:id>')
@login_required
@roles_required('Administrador', 'Secretaria')
def confirmar_cita(id):
    cita = Cita.query.get_or_404(id)
    cita.estado = 'Confirmada'
    db.session.commit()
    if cita.id_bot_sic:
        fecha = cita.fecha_cita.strftime('%d/%m/%Y')
        hora = cita.fecha_cita.strftime('%H:%M')
        notificar_paciente(cita.id_bot_sic, f"✅ *¡Tu cita ha sido CONFIRMADA!*\n📅 Fecha: {fecha}\n⏰ Hora: {hora}\nTe esperamos para tu consulta.🏥")
    flash("Cita confirmada y paciente notificado.", "success")
    return redirect(url_for('citas'))


@app.route('/citas/rechazar/<int:id>')
@login_required
@roles_required('Administrador', 'Secretaria')
def rechazar_cita(id):
    cita = Cita.query.get_or_404(id)
    cita.estado = 'Cancelada'
    db.session.commit()
    if cita.id_bot_sic:
        fecha = cita.fecha_cita.strftime('%d/%m/%Y')
        hora = cita.fecha_cita.strftime('%H:%M')
        notificar_paciente(cita.id_bot_sic, f"❌ *Tu cita del {fecha} a las {hora} ha sido cancelada.*\nLamentamos los inconvenientes. Contáctanos para reprogramar.")
    flash("Cita cancelada y paciente notificado.", "success")
    return redirect(url_for('citas'))


@app.route('/usuarios')
@login_required
@roles_required('Administrador')
def lista_usuarios():
    usuarios = db.session.query(Login, Usuario, Rol)\
        .join(Usuario, Login.id_usuario == Usuario.id_usuario)\
        .join(Rol, Login.id_rol == Rol.id_rol)\
        .all()
    roles = Rol.query.all()
    return render_template('usuarios.html', usuarios=usuarios, roles=roles)


@app.route('/usuarios/bloquear/<int:id>')
@login_required
@roles_required('Administrador')
def bloquear_usuario(id):
    u = Login.query.get_or_404(id)
    if u.id_login == current_user.id_login:
        flash("No puedes bloquearte a ti mismo.", "error")
        return redirect(url_for('lista_usuarios'))
    u.activo = not u.activo
    db.session.commit()
    estado = "desbloqueado" if u.activo else "bloqueado"
    flash(f"Usuario '{u.usuario}' {estado} correctamente.", "success")
    return redirect(url_for('lista_usuarios'))


@app.route('/usuarios/eliminar/<int:id>')
@login_required
@roles_required('Administrador')
def eliminar_usuario(id):
    u = Login.query.get_or_404(id)
    if u.id_login == current_user.id_login:
        flash("No puedes eliminarte a ti mismo.", "error")
        return redirect(url_for('lista_usuarios'))
    try:
        id_usuario = u.id_usuario
        # 1. Borrar sesiones activas del usuario (FK constraint)
        with db.engine.connect() as conn:
            conn.execute(db.text("DELETE FROM sessiones WHERE id_usuario = :uid"), {"uid": id_usuario})
            conn.commit()
        # 2. Borrar login
        db.session.delete(u)
        db.session.flush()
        # 3. Borrar perfil de usuario
        perfil = Usuario.query.get(id_usuario)
        if perfil:
            db.session.delete(perfil)
        db.session.commit()
        flash("Usuario eliminado correctamente.", "success")
    except Exception as e:
        db.session.rollback()
        flash(f"Error al eliminar: {str(e)}", "error")
    return redirect(url_for('lista_usuarios'))



@app.route('/usuarios/editar/<int:id>', methods=['POST'])
@login_required
@roles_required('Administrador')
def editar_usuario(id):
    u = Login.query.get_or_404(id)
    perfil = Usuario.query.get(u.id_usuario)
    nuevo_usuario = request.form.get('usuario')
    nueva_clave = request.form.get('password')
    nuevo_rol = request.form.get('id_rol')
    if nuevo_usuario:
        u.usuario = nuevo_usuario
    if nueva_clave:
        u.contrasena = nueva_clave
    if nuevo_rol:
        u.id_rol = nuevo_rol
    if perfil:
        perfil.nombres = request.form.get('nombres', perfil.nombres)
        perfil.apellidos = request.form.get('apellidos', perfil.apellidos)
        perfil.correo_electronico = request.form.get('correo', perfil.correo_electronico)
    db.session.commit()
    flash(f"Usuario '{u.usuario}' actualizado correctamente.", "success")
    return redirect(url_for('lista_usuarios'))


# ═══════════════════════════════════════════════════════════
# ARRANQUE DE LA APLICACIÓN
# ═══════════════════════════════════════════════════════════
def start_flask():
    app.run(host='0.0.0.0', port=5000, debug=True, use_reloader=False)

class PywebviewAPI:
    def guardar_excel_nativo(self, sugerencia_nombre, data_base64):
        import webview
        import base64
        import traceback
        try:
            window = webview.windows[0]
            if not window:
                return {"status": "error", "message": "No main window"}
            
            # Pedir al usuario donde guardar
            resultado = window.create_file_dialog(
                webview.SAVE_DIALOG,
                directory='',
                save_filename=sugerencia_nombre,
                file_types=('Excel Files (*.xlsx)', 'All files (*.*)')
            )
            
            if resultado and len(resultado) > 0:
                ruta_archivo = resultado[0]
                with open(ruta_archivo, 'wb') as f:
                    f.write(base64.b64decode(data_base64))
                return {"status": "success"}
            return {"status": "cancelled"}
        except Exception as e:
            print("Error guardando nativo:", str(e))
            traceback.print_exc()
            return {"status": "error", "message": str(e)}


if __name__ == '__main__':
    # Intentar obtener la tasa BCV al arrancar
    try:
        obtener_tasa_bcv()
    except:
        pass

    api = PywebviewAPI()

    t = Thread(target=start_flask)
    t.daemon = True
    t.start()

    webview.create_window('Sistema SIC - Clinica', 'http://127.0.0.1:5000', width=1200, height=800, js_api=api)
    webview.start(icon=resource_path('static/logo.png'))